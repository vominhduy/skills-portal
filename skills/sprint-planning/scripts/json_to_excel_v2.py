#!/usr/bin/env python3
"""
json_to_excel_v2.py - Convert sprint backlog JSON to Detailed Sprint Backlog Excel.
Template: Sprint Backlog with User Stories, Story Points, Manday, Task Breakdown, Status, PIC, Dates.
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation


# ====================== CẤU HÌNH ======================
def parse_args():
    parser = argparse.ArgumentParser(description="Convert sprint backlog JSON to Detailed Sprint Backlog Excel")
    parser.add_argument("json_file", nargs="?", default="sprint1_backlog.json",
                        help="Path to sprint backlog JSON file")
    parser.add_argument("output_file", nargs="?", default=None,
                        help="Optional positional output Excel file path")
    parser.add_argument("-o", "--output", default=None,
                        help="Output Excel file path (default: Sprint_Backlog_Raw.xlsx)")
    parser.add_argument("--output-dir", default=None,
                        help="Output directory injected by webapp")
    return parser.parse_known_args()[0]


_shared_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "_shared"))
if _shared_dir not in sys.path:
    sys.path.insert(0, _shared_dir)
from output_path import ensure_versioned_path, resolve_input_path

args = parse_args()
JSON_FILE = resolve_input_path(args.json_file)

# Tự động tạo tên sheet từ tên file (sprint1_backlog.json -> Sprint 1)
_match = re.search(r"sprint(\d+)", JSON_FILE, re.IGNORECASE)
_sprint_num = _match.group(1) if _match else "1"
NEW_SHEET_NAME = f"Sprint {_sprint_num}"

# Tên file Excel sẽ được tạo (bắt buộc kèm _v{N}_{HHMMSS})
_raw_excel_file = args.output or args.output_file or "Sprint_Backlog_Raw.xlsx"
if args.output_dir and not os.path.isabs(_raw_excel_file):
    _raw_excel_file = os.path.join(args.output_dir, os.path.basename(_raw_excel_file))
EXCEL_FILE = ensure_versioned_path(_raw_excel_file)


def optimize_for_excel_online(file_path):
    try:
        _shared_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "_shared"))
        if _shared_dir not in sys.path:
            sys.path.insert(0, _shared_dir)
        from xlsx_normalizer import normalize_xlsx_file
        normalize_xlsx_file(file_path)
    except Exception:
        pass

# ==================== MÀU SẮC ====================
COLOR_HEADER = PatternFill(start_color="D6DCE4", end_color="D6DCE4", fill_type="solid")
COLOR_EPIC = PatternFill(start_color="FFE598", end_color="FFE598", fill_type="solid")
COLOR_DONE = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
COLOR_INPROGRESS = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")
COLOR_PENDING = PatternFill(start_color="A6A6A6", end_color="A6A6A6", fill_type="solid")
COLOR_NOTSTARTED = PatternFill(start_color="EFF6FB", end_color="EFF6FB", fill_type="solid")
COLOR_CRITICAL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
COLOR_HIGH = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
COLOR_MEDIUM = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
COLOR_LOW = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
COLOR_DARK_HEADER = PatternFill(start_color="305496", end_color="305496", fill_type="solid")

# ==================== BORDER ====================
thin = Side(style="thin")
dashed = Side(style="dashed")
thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
dashed_border = Border(left=dashed, right=dashed, top=dashed, bottom=dashed)

DATE_FORMAT = "[$-en-US]ddd dd/MMM/yy"

# ==================== HÀM CLEAN DASH (từ đầu JSON) ====================
def clean_text(text):
    if isinstance(text, str):
        return (text
                .replace("—", "-")   # em dash
                .replace("–", "-")   # en dash
                .replace("‐", "-")   # hyphen khác
                .replace("‑", "-"))
    return text

def clean_json_data(obj):
    """Clean recursive toàn bộ JSON (dict, list, string)"""
    if isinstance(obj, str):
        return clean_text(obj)
    elif isinstance(obj, list):
        return [clean_json_data(item) for item in obj]
    elif isinstance(obj, dict):
        return {key: clean_json_data(value) for key, value in obj.items()}
    else:
        return obj


def parse_date(date_str):
    if not date_str or str(date_str).strip() == "":
        return None
    try:
        return datetime.strptime(str(date_str).strip(), "%d/%m/%Y")
    except:
        return None


def load_backlog_data():
    if not os.path.exists(JSON_FILE):
        print(f"❌ Không tìm thấy file JSON: {JSON_FILE}")
        return None
    with open(JSON_FILE, "r", encoding="utf-8") as f:
        raw_json = f.read()                    # Đọc thành chuỗi thô
        # Clean toàn bộ chuỗi JSON trước khi parse
        cleaned_json = clean_text(raw_json)    # Thay thế tất cả dash trong toàn bộ file
        data = json.loads(cleaned_json)        # Parse JSON đã sạch

    # Clean thêm một lần recursive (đảm bảo 100% sạch)
    data = clean_json_data(data)
    return data


def validate_task_effort(data, max_hours=24):
    """Cảnh báo các task vi phạm giới hạn est_effort_hour > max_hours."""
    violations = []
    for epic in data.get("epics", []):
        for us in epic.get("user_stories", []):
            for task in us.get("tasks", []):
                hours = task.get("est_effort_hour", 0)
                try:
                    hours = float(hours)
                except (TypeError, ValueError):
                    hours = 0
                if hours > max_hours:
                    violations.append({
                        "us_id": us.get("us_id", "?"),
                        "pic": task.get("pic", "?"),
                        "task_name": task.get("task_name", "?")[:60],
                        "hours": hours
                    })
    if violations:
        print(f"\n⚠️  CẢNH BÁO: {len(violations)} task vượt quá {max_hours}h - nên tách nhỏ hơn:")
        for v in violations:
            print(f"   [{v['us_id']}] {v['pic']} | {v['hours']}h | {v['task_name']}")
        print()
    return violations


def create_new_sprint_excel():
    data = load_backlog_data()
    if not data:
        return

    # Kiểm tra task vượt quá 24h
    validate_task_effort(data)

    # Tạo Workbook mới hoàn toàn
    wb = Workbook()
    ws = wb.active
    ws.title = NEW_SHEET_NAME

    # ==================== HEADER ====================
    ws["B1"] = f"SPRINT BACKLOG #{_sprint_num}"
    ws["B1"].font = Font(bold=True, size=28, color="ED7D31")
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center")

    ws["B3"] = "SPRINT GOAL"
    ws["B3"].fill = COLOR_DARK_HEADER
    ws["B3"].font = Font(bold=True, color="FFFFFF")
    ws["B3"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("C3:G3")
    ws["C3"] = data.get("sprint_goal", "")
    ws["C3"].font = Font(color="C65911")
    ws["C3"].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

    ws["B4"] = "TIME BOX"
    ws["B4"].fill = COLOR_DARK_HEADER
    ws["B4"].font = Font(bold=True, color="FFFFFF")
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("C4:G4")
    time_box = data.get("time_box", "")
    time_box_text = "\n".join(str(item).strip() for item in time_box) if isinstance(time_box, list) else str(time_box).replace("\\n", "\n")
    ws["C4"] = time_box_text
    ws["C4"].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

    ws.row_dimensions[3].height = 30
    ws.row_dimensions[4].height = 75

    ws["H3"] = "STORY POINT OF THIS SPRINT"
    ws["H3"].fill = COLOR_DARK_HEADER
    ws["H3"].font = Font(bold=True, color="FFFFFF")
    ws["H3"].alignment = Alignment(horizontal="center", vertical="center")

    ws["I3"] = "MANDAY"
    ws["I3"].fill = COLOR_DARK_HEADER
    ws["I3"].font = Font(bold=True, color="FFFFFF", size=12)
    ws["I3"].alignment = Alignment(horizontal="center", vertical="center")

    for row in range(3, 5):
        for col in range(2, 10):
            ws.cell(row=row, column=col).border = dashed_border

    # ==================== TABLE HEADER ====================
    start_row = 6
    headers = ["User\nStory\nID", "MODULE/ FEATURE", "USER STORY/DESCRIPTION",
               "STORY\nPOINT", "ESTIMATED\nEFFORT\n(Story\nManday)", "Priority",
               "Complexity", "TASK", "PIC", "ESTIMATED\nEFFORT\n(Task\nManday)",
               "STATUS", "START DATE", "DUE DATE", "WORKING\nDAY", "ESTIMATED\nEFFORT\n(HOUR)",
               "NOTES",
               "TASK COMPLEXITY", "Type OF WORK", "TASK PROCESS", "TASK PRIORITY"]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True)
        if col > 10:
            cell.fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        else:
            cell.fill = COLOR_HEADER

        if col == 10 or col == 14:
            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            cell.font = Font(bold=True, color="000000")

        if col > 15:
            cell.font = Font(bold=True, color="FFD965")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # ==================== GHI DỮ LIỆU ====================
    current_row = start_row + 1
    first_data_row = current_row
    first_epic_row = None
    epic_rows = []

    for epic in data.get("epics", []):
        epic_row = current_row
        if first_epic_row is None:
            first_epic_row = epic_row

        epic_rows.append(epic_row)

        # Dòng Epic
        epic_id_cell = ws.cell(row=epic_row, column=2, value=epic.get("epic", ""))
        epic_id_cell.font = Font(bold=True)
        epic_id_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

        for c in range(1, 21):
            cell = ws.cell(row=epic_row, column=c)
            cell.fill = COLOR_EPIC
            cell.font = Font(bold=True)
            cell.border = dashed_border

            if c == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            elif c in [12, 13]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.cell(row=epic_row, column=1).value = None
        current_row += 1

        epic_us_start_row = current_row
        epic_tasks_start = None

        for us in epic.get("user_stories", []):
            us_start_row = current_row

            # User Story
            ws.cell(row=us_start_row, column=1, value=us.get("us_id"))
            ws.cell(row=us_start_row, column=2, value=us.get("module"))
            ws.cell(row=us_start_row, column=3, value=us.get("description"))
            ws.cell(row=us_start_row, column=4).value = f"=AVERAGE(E{us_start_row},(2*LEFT(F{us_start_row},1)),(2*LEFT(G{us_start_row},1)))"
            ws.cell(row=us_start_row, column=4).number_format = "0.0"
            ws.cell(row=us_start_row, column=5).number_format = "0.0"
            ws.cell(row=us_start_row, column=6, value=us.get("priority", ""))
            ws.cell(row=us_start_row, column=7, value=us.get("complexity", ""))

            tasks = us.get("tasks", [])
            if tasks:
                if epic_tasks_start is None:
                    epic_tasks_start = current_row

                first_task = tasks[0]
                ws.cell(row=us_start_row, column=8, value=first_task.get("task_name"))
                ws.cell(row=us_start_row, column=9, value=first_task.get("pic"))
                ws.cell(row=us_start_row, column=10).value = f"=O{us_start_row}/8"
                ws.cell(row=us_start_row, column=11, value=first_task.get("status"))
                ws.cell(row=us_start_row, column=12, value=parse_date(first_task.get("start_date")))
                ws.cell(row=us_start_row, column=13, value=parse_date(first_task.get("due_date")))
                ws.cell(row=us_start_row, column=12).number_format = DATE_FORMAT
                ws.cell(row=us_start_row, column=13).number_format = DATE_FORMAT
                ws.cell(row=us_start_row, column=14).value = f"=NETWORKDAYS(L{us_start_row},M{us_start_row})"
                ws.cell(row=us_start_row, column=15, value=float(first_task.get("est_effort_hour", 0)))
                ws.cell(row=us_start_row, column=16, value=first_task.get("notes"))
                ws.cell(row=us_start_row, column=17, value=first_task.get("task_complexity"))
                ws.cell(row=us_start_row, column=18, value=first_task.get("type_of_work"))
                ws.cell(row=us_start_row, column=19, value=first_task.get("task_process"))
                ws.cell(row=us_start_row, column=20, value=first_task.get("task_priority"))

                for task in tasks[1:]:
                    current_row += 1
                    ws.cell(row=current_row, column=8, value=task.get("task_name"))
                    ws.cell(row=current_row, column=9, value=task.get("pic"))
                    ws.cell(row=current_row, column=10).value = f"=O{current_row}/8"
                    ws.cell(row=current_row, column=11, value=task.get("status"))
                    ws.cell(row=current_row, column=12, value=parse_date(task.get("start_date")))
                    ws.cell(row=current_row, column=13, value=parse_date(task.get("due_date")))
                    ws.cell(row=current_row, column=12).number_format = DATE_FORMAT
                    ws.cell(row=current_row, column=13).number_format = DATE_FORMAT
                    ws.cell(row=current_row, column=14).value = f"=NETWORKDAYS(L{current_row},M{current_row})"
                    ws.cell(row=current_row, column=15, value=float(task.get("est_effort_hour", 0)))
                    ws.cell(row=current_row, column=16, value=task.get("notes"))
                    ws.cell(row=current_row, column=17, value=task.get("task_complexity"))
                    ws.cell(row=current_row, column=18, value=task.get("type_of_work"))
                    ws.cell(row=current_row, column=19, value=task.get("task_process"))
                    ws.cell(row=current_row, column=20, value=task.get("task_priority"))

            us_end_row = current_row

            # Merge cột A-G cho User Story
            if us_end_row > us_start_row:
                for col in range(1, 8):
                    ws.merge_cells(start_row=us_start_row, start_column=col, end_row=us_end_row, end_column=col)

            # Alignment User Story
            for col in range(1, 8):
                ws.cell(row=us_start_row, column=col).alignment = Alignment(vertical="center", wrap_text=True)

            if tasks:
                ws.cell(row=us_start_row, column=5).value = f"=SUM(J{us_start_row}:J{us_end_row})"

            current_row += 1

        # Tính toán dòng Epic
        epic_end_row = current_row - 1
        if epic_us_start_row <= epic_end_row:
            ws.cell(row=epic_row, column=4).value = f"=SUM(D{epic_us_start_row}:D{epic_end_row})"
            ws.cell(row=epic_row, column=4).number_format = "0.0"
            ws.cell(row=epic_row, column=5).value = f"=SUM(E{epic_us_start_row}:E{epic_end_row})"
            ws.cell(row=epic_row, column=5).number_format = "0.0"

        if epic_tasks_start and epic_tasks_start <= epic_end_row:
            ws.cell(row=epic_row, column=15).value = f"=SUM(O{epic_tasks_start}:O{epic_end_row})"
            ws.cell(row=epic_row, column=15).font = Font(bold=True)
            ws.cell(row=epic_row, column=10).value = f"=O{epic_row}/8"
            ws.cell(row=epic_row, column=10).font = Font(bold=True)
            ws.cell(row=epic_row, column=12).value = f"=MIN(L{epic_tasks_start}:L{epic_end_row})"
            ws.cell(row=epic_row, column=13).value = f"=MAX(M{epic_tasks_start}:M{epic_end_row})"
            ws.cell(row=epic_row, column=12).number_format = DATE_FORMAT
            ws.cell(row=epic_row, column=13).number_format = DATE_FORMAT
            ws.cell(row=epic_row, column=14).value = f"=NETWORKDAYS(L{epic_row},M{epic_row})"
            ws.cell(row=epic_row, column=14).font = Font(bold=True)

    # ==================== TỔNG STORY POINT & MANDAY ====================
    last_row = current_row - 1
    if first_epic_row:
        ws["I4"].value = f"=SUM(O{first_epic_row}:O{last_row})/8/2"
        ws["I4"].font = Font(bold=True, size=26, color="ED7D31")
        ws["I4"].alignment = Alignment(horizontal="center", vertical="center")
        ws["I4"].number_format = "0.0"

        ws["H4"].value = f"=SUM(D{first_data_row}:D{last_row})/2"
        ws["H4"].font = Font(bold=True, size=26, color="ED7D31")
        ws["H4"].alignment = Alignment(horizontal="center", vertical="center")
        ws["H4"].number_format = "0.0"

    # ==================== ALIGNMENT & BORDER ====================
    for row in range(start_row + 1, current_row):
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=row, column=8).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        # Căn giữa cho cột I (PIC)
        ws.cell(row=row, column=9).alignment = Alignment(horizontal="center", vertical="center")

    for epic_row in epic_rows:
        ws.cell(row=epic_row, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

    center_align = Alignment(horizontal="center", vertical="center")
    for row in range(start_row + 1, current_row):
        for col in [4, 5, 6, 7, 10, 11, 12, 13, 14, 15]:
            cell = ws.cell(row=row, column=col)
            if cell.value is not None or col in [12, 13]:
                cell.alignment = center_align

    # Border
    for row in range(start_row, current_row):
        for col in range(1, 21):
            ws.cell(row=row, column=col).border = dashed_border

    # Wrap text cho cột Notes (P)
    for row in range(start_row + 1, current_row):
        cell = ws.cell(row=row, column=16)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ==================== DATA VALIDATION & CONDITIONAL FORMATTING ====================
    max_row = current_row + 50

    dv_priority = DataValidation(type="list", formula1='"4-Critical,3-High,2-Medium,1-Low"', allow_blank=True)
    ws.add_data_validation(dv_priority)
    dv_priority.add(f"F{start_row+1}:F{max_row}")

    dv_complexity = DataValidation(type="list", formula1='"3-High,2-Medium,1-Low"', allow_blank=True)
    ws.add_data_validation(dv_complexity)
    dv_complexity.add(f"G{start_row+1}:G{max_row}")

    dv_status = DataValidation(type="list", formula1='"Done,In Progress,Not Started,Pending"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add(f"K{start_row+1}:K{max_row}")

    # Conditional Formatting
    cf_p = f"F{start_row+1}:F{max_row}"
    ws.conditional_formatting.add(cf_p, CellIsRule(operator="equal", formula=['"4-Critical"'], fill=COLOR_CRITICAL, font=Font(color="FFFFFF")))
    ws.conditional_formatting.add(cf_p, CellIsRule(operator="equal", formula=['"3-High"'], fill=COLOR_HIGH, font=Font(color="9C0006")))
    ws.conditional_formatting.add(cf_p, CellIsRule(operator="equal", formula=['"2-Medium"'], fill=COLOR_MEDIUM, font=Font(color="9C5700")))
    ws.conditional_formatting.add(cf_p, CellIsRule(operator="equal", formula=['"1-Low"'], fill=COLOR_LOW, font=Font(color="1F497D")))

    cf_c = f"G{start_row+1}:G{max_row}"
    ws.conditional_formatting.add(cf_c, CellIsRule(operator="equal", formula=['"3-High"'], fill=COLOR_HIGH, font=Font(color="9C0006")))
    ws.conditional_formatting.add(cf_c, CellIsRule(operator="equal", formula=['"2-Medium"'], fill=COLOR_MEDIUM, font=Font(color="9C5700")))
    ws.conditional_formatting.add(cf_c, CellIsRule(operator="equal", formula=['"1-Low"'], fill=COLOR_LOW, font=Font(color="1F497D")))

    cf_k = f"K{start_row+1}:K{max_row}"
    ws.conditional_formatting.add(cf_k, CellIsRule(operator="equal", formula=['"Done"'], fill=COLOR_DONE, font=Font(color="006100")))
    ws.conditional_formatting.add(cf_k, CellIsRule(operator="equal", formula=['"In Progress"'], fill=COLOR_INPROGRESS, font=Font(color="D55709")))
    ws.conditional_formatting.add(cf_k, CellIsRule(operator="equal", formula=['"Pending"'], fill=COLOR_PENDING, font=Font(color="262626")))

    ws.freeze_panes = "H7"

    # Column Widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 10
    ws.column_dimensions["H"].width = 50
    ws.column_dimensions["I"].width = 16
    ws.column_dimensions["J"].width = 12
    ws.column_dimensions["K"].width = 12
    ws.column_dimensions["L"].width = 15
    ws.column_dimensions["M"].width = 15
    ws.column_dimensions["N"].width = 10
    ws.column_dimensions["O"].width = 12
    ws.column_dimensions["P"].width = 50
    ws.column_dimensions["Q"].width = 10
    ws.column_dimensions["R"].width = 10
    ws.column_dimensions["S"].width = 10
    ws.column_dimensions["T"].width = 10

    # Lưu file Excel mới
    wb.save(EXCEL_FILE)
    try:
        _script_dir = os.path.dirname(os.path.abspath(__file__))
        if _script_dir not in sys.path:
            sys.path.insert(0, _script_dir)
        from generate_xlsxwriter import build_sprint_backlog
        build_sprint_backlog(data, EXCEL_FILE)
    except Exception as e:
        print(f"Warning: XlsxWriter fallback error: {e}")
    optimize_for_excel_online(EXCEL_FILE)
    print(f"✅ Đã tạo thành công file Excel Sprint Backlog: **{EXCEL_FILE}**")
    print(f" -> Sheet: '{NEW_SHEET_NAME}'")
    print(f" -> Số dòng Epic: {len(epic_rows)}")


if __name__ == "__main__":
    create_new_sprint_excel()
