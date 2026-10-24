"""Fetch sprint data from Excel files (.xlsx) or Microsoft SharePoint / OneDrive (Graph API)
and output standardized JSON.

Supports 3 formats:
  - sprint_backlog: Sprint Backlog template (US + Task hierarchy)
  - jira:           Jira export format
  - generic:        Any other header layout, best-effort column matching

Nguồn dữ liệu:
  - File Excel cục bộ (.xlsx) qua thư viện openpyxl.
  - File Excel Online trên SharePoint / OneDrive qua Microsoft Graph API .
"""
import argparse
import json
import os
import sys
import re
from datetime import datetime, timezone, timedelta, date
from typing import Dict, List, Optional, Any

_HERE = os.path.dirname(os.path.abspath(__file__))
_SHARED_DIR = os.path.join(_HERE, "..", "..", "_shared")
if _SHARED_DIR not in sys.path:
    sys.path.insert(0, _SHARED_DIR)

try:
    import ms_graph_client as mg
except ImportError:
    mg = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

VN_TZ = timezone(timedelta(hours=7))

# ──────────────────── DYNAMIC HEADER MATCHING ──────────────────

SPRINT_FIELD_PATTERNS = {
    "us_id":           ["user story id", "story id", "us id", "id"],
    "module":          ["module/feature", "module/ feat", "module", "feature"],
    "description":     ["user story/description", "user story desc", "description",
                        "user story", "story description"],
    "story_point":     ["story point", "point", "sp"],
    "priority":        ["priority"],
    "task_name":       ["task"],
    "pic":             ["pic", "assignee", "person in charge", "person"],
    "status":          ["status"],
    "start_date":      ["start date", "start"],
    "due_date":        ["due date", "due", "end date"],
    "working_day":     ["working day", "work day", "workingdays", "workdays"],
    "effort_hour":     ["estimated effort (hour)", "estimated effort hour",
                        "effort (hour)", "effort hour", "hour", "man-hour"],
    "notes":           ["notes", "note", "comment"],
    "task_process":    ["task process", "process"],
    "task_priority":   ["task priority"],
    "complexity":      ["complexity", "task complexity"],
    "type_of_work":    ["type of work", "work type"],
    "estimated_manday": ["estimated effort (story manday)", "estimated effort",
                         "story manday", "manday", "man-day"],
    "task_manday":     ["estimated effort (task manday)", "task manday"],
}

JIRA_FIELD_PATTERNS = {
    "summary":          ["summary"],
    "id":               ["issue key", "key", "issue id", "id"],
    "issue_type":       ["issue type", "type"],
    "status":           ["status"],
    "priority":         ["priority"],
    "assignee":         ["assignee", "assignee name"],
    "due_date":         ["due date", "due"],
    "description":      ["description", "desc"],
    "original_estimate": ["original estimate", "original estimate", "estimate"],
    "remaining_estimate": ["remaining estimate"],
    "time_spent":       ["time spent", "time spent"],
    "work_ratio":       ["work ratio"],
    "parent_key":       ["parent key", "parent"],
    "start_date":       ["start date", "start"],
    "task_process":     ["task process", "process"],
    "complexity":       ["complexity"],
    "resolution":       ["resolution"],
}

GENERIC_FIELD_PATTERNS = {
    "id":               ["issue key", "us id", "story id", "task id", "key", "id"],
    "summary":          ["summary", "task name", "user story", "description", "task", "title"],
    "status":           ["status", "state"],
    "assignee":         ["assignee", "pic", "person in charge", "owner"],
    "priority":         ["priority"],
    "start_date":       ["start date", "start"],
    "due_date":         ["due date", "due", "end date", "deadline"],
    "estimate":         ["original estimate", "estimate", "effort (hour)", "effort hour",
                         "effort", "story point", "man-day", "manday", "hour"],
    "time_spent":       ["time spent"],
    "parent_key":       ["parent key", "parent"],
}


def build_column_index(headers, field_patterns):
    """Xây dựng index cột động: field_name → column_index."""
    normalized = {}
    for i, h in enumerate(headers):
        if h:
            normalized[i] = re.sub(r'\s+', ' ', str(h).lower().replace("\n", " ").replace("\r", " ")).strip()
    
    col_idx = {}
    used = set()
    
    for field, patterns in field_patterns.items():
        best_score, best_col = 0, None
        for col, text in normalized.items():
            if col in used:
                continue
            for p in patterns:
                if p in text and len(p) > best_score:
                    best_score = len(p)
                    best_col = col
        if best_col is not None:
            col_idx[field] = best_col
            used.add(best_col)
    
    if col_idx:
        print(f"  → Column index: {dict(sorted(col_idx.items(), key=lambda x: x[1]))}")
    else:
        print("  → Cảnh báo: Không khớp cột nào! Dùng tất cả headers dạng thô.")
    return col_idx


SPRINT_KEYS = {
    "user story id", "module/feature", "module/ feature", "user story/description",
    "story point", "estimated effort", "task", "pic", "status",
    "start date", "due date", "working day", "task process", "task priority",
    "complexity", "notes", "type of work",
}

JIRA_KEYS = {
    "summary", "issue key", "issue type", "status", "priority",
    "resolution", "assignee", "due date", "description",
    "original estimate", "remaining estimate", "time spent",
    "work ratio", "parent key",
}


def detect_format(headers):
    """Nhận diện định dạng bảng tính."""
    header_lower = set()
    for h in headers:
        if h:
            normalized = str(h).lower().replace("\n", " ").strip()
            normalized = re.sub(r'\s+', ' ', normalized)
            header_lower.add(normalized)

    sprint_matches = sum(1 for kw in SPRINT_KEYS for hl in header_lower if kw in hl)
    jira_matches = sum(1 for kw in JIRA_KEYS for hl in header_lower if kw in hl)

    print(f"  → Format scores - sprint_backlog:{sprint_matches}  jira:{jira_matches}")

    if sprint_matches >= jira_matches and sprint_matches >= 6:
        return "sprint_backlog"
    elif jira_matches >= 5:
        return "jira"
    return "generic"


def parse_date_from_str(date_str):
    if not date_str or not isinstance(date_str, str):
        return None
    date_str = date_str.strip()
    if not date_str:
        return None
    formats = [
        "%d/%b/%y %I:%M %p",
        "%d-%b-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d %b %Y",
        "%m/%d/%Y",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(date_str, fmt).date()
        except ValueError:
            continue
    return None


def parse_date_from_cell(cell_value):
    if cell_value is None:
        return None
    if isinstance(cell_value, datetime):
        return cell_value.date()
    if isinstance(cell_value, date):
        return cell_value
    return parse_date_from_str(str(cell_value))


def str_cell(val):
    if val is None:
        return ""
    return str(val).strip()


def safe_float(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", "").replace(" ", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def parse_sprint_backlog(rows, headers):
    col = build_column_index(headers, SPRINT_FIELD_PATTERNS)
    if not col:
        print("  ⚠️ Không detect được cột, dùng raw headers")
        return parse_raw(rows, headers)
    
    stories = []
    current_story = None
    
    def cv(row, field, default=""):
        c = col.get(field)
        if c is not None and c < len(row):
            return str_cell(row[c])
        return default
    
    def cvn(row, field, default=0.0):
        c = col.get(field)
        if c is not None and c < len(row):
            return safe_float(row[c])
        return default
    
    for row_idx, row in enumerate(rows):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        
        us_id = cv(row, "us_id")
        module = cv(row, "module")
        desc = cv(row, "description")
        task_name = cv(row, "task_name")
        
        is_epic = (not us_id and module.upper().startswith("EPIC"))
        is_us = bool(us_id)
        
        if is_epic:
            print(f"  [EPIC] Row {row_idx+8}: {module[:60]}")
            current_story = None
            continue
        
        if is_us:
            priority_val = cv(row, "priority", "2-Medium")
            story_point = cvn(row, "story_point")
            effort_hour = cvn(row, "effort_hour")
            
            current_story = {
                "id": us_id,
                "summary": desc[:200] if desc else f"{module} - {task_name}",
                "status": "To Do",
                "assignee": "Unassigned",
                "priority": priority_val,
                "issue_type": "Story",
                "start_date": None,
                "due_date": None,
                "task_process": "",
                "original_estimate_h": effort_hour if effort_hour > 0 else story_point * 8,
                "time_spent_h": 0,
                "remaining_estimate_h": 0,
                "parent_key": "",
                "pct_done": 0,
                "module": module,
                "subtasks": [],
            }
            
            if task_name:
                subtask = _parse_sb_task(row, us_id, col)
                if subtask:
                    current_story["subtasks"].append(subtask)
            
            stories.append(current_story)
            print(f"  [US]  {us_id}: {desc[:50] if desc else task_name[:50]}... ({len(current_story['subtasks'])} tasks)")
        
        elif task_name and current_story is not None:
            subtask = _parse_sb_task(row, current_story["id"], col)
            if subtask:
                current_story["subtasks"].append(subtask)
    
    print(f"\n✅ Sprint Backlog: {len(stories)} stories, {sum(len(s['subtasks']) for s in stories)} subtasks")
    return stories, [], []


def _parse_sb_task(row, parent_us_id, col):
    def cv(field, default=""):
        c = col.get(field)
        return str_cell(row[c]) if c is not None and c < len(row) else default
    
    def cvn(field, default=0.0):
        c = col.get(field)
        return safe_float(row[c]) if c is not None and c < len(row) else default
    
    task_name = cv("task_name")
    if not task_name:
        return None
    
    pic = cv("pic", "Unassigned")
    status = cv("status", "To Do")
    effort_hour = cvn("effort_hour")
    task_process = cv("task_process")
    task_priority = cv("task_priority", "Medium")
    notes = cv("notes")
    working_day = cvn("working_day")
    
    c = col.get("start_date")
    start_dt = parse_date_from_cell(row[c]) if c is not None and c < len(row) else None
    c = col.get("due_date")
    due_dt = parse_date_from_cell(row[c]) if c is not None and c < len(row) else None
    
    status_lower = status.lower()
    if status_lower in ("done", "completed", "closed", "resolved"):
        pct_done = 100
    elif status_lower in ("in progress", "doing", "inreview", "in review"):
        pct_done = 50
    else:
        pct_done = 0
    
    if effort_hour == 0 and working_day > 0:
        effort_hour = working_day * 8
    
    return {
        "id": f"{parent_us_id}_{_slugify(task_name[:20])}",
        "summary": task_name[:200],
        "status": status,
        "assignee": pic,
        "priority": task_priority or "Medium",
        "issue_type": "Sub-task",
        "start_date": start_dt,
        "due_date": due_dt,
        "task_process": task_process,
        "original_estimate_h": effort_hour,
        "time_spent_h": 0,
        "remaining_estimate_h": effort_hour,
        "parent_key": parent_us_id,
        "pct_done": pct_done,
        "notes": notes,
        "working_day": working_day,
    }


def _slugify(text):
    slug = re.sub(r'[^a-zA-Z0-9]+', '_', text.lower()).strip('_')
    return slug[:30] if slug else "task"


def parse_jira(rows, headers):
    col = build_column_index(headers, JIRA_FIELD_PATTERNS)
    if not col:
        print("  ⚠️ Không detect được cột Jira, dùng raw headers")
        return parse_raw(rows, headers)
    
    stories, subtasks, bugs = [], [], []
    
    for row in rows:
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        
        def cv(field, default=""):
            c = col.get(field)
            return str_cell(row[c]) if c is not None and c < len(row) else default
        
        def cvn(field, default=0.0):
            c = col.get(field)
            return safe_float(row[c]) if c is not None and c < len(row) else default
        
        issue_id = cv("id")
        issue_type = cv("issue_type")
        parent_key = cv("parent_key")
        status = cv("status", "To Do") or "To Do"
        summary = cv("summary")
        assignee = cv("assignee", "Unassigned") or "Unassigned"
        priority = cv("priority", "Medium") or "Medium"
        
        c = col.get("start_date")
        start_date = parse_date_from_cell(row[c]) if c is not None and c < len(row) else None
        c = col.get("due_date")
        due_date = parse_date_from_cell(row[c]) if c is not None and c < len(row) else None
        task_process = cv("task_process")
        
        original_estimate = parse_estimate_to_hours(cvn("original_estimate"))
        time_spent = parse_estimate_to_hours(cvn("time_spent"))
        remaining_estimate = parse_estimate_to_hours(cvn("remaining_estimate"))
        pct_done = parse_pct_done(cv("work_ratio", "0%"))
        
        task = {
            "id": issue_id,
            "summary": summary[:200],
            "status": status,
            "assignee": assignee,
            "priority": priority,
            "issue_type": issue_type,
            "start_date": start_date,
            "due_date": due_date,
            "task_process": task_process,
            "original_estimate_h": original_estimate,
            "time_spent_h": time_spent,
            "remaining_estimate_h": remaining_estimate,
            "parent_key": parent_key,
            "pct_done": pct_done,
        }
        
        if issue_type == "Story":
            stories.append(task)
        elif issue_type in ("Sub-task", "Subtask"):
            subtasks.append(task)
        elif issue_type == "Bug":
            bugs.append(task)
    
    print(f"\n✅ Jira format: {len(stories)} stories, {len(subtasks)} subtasks, {len(bugs)} bugs")
    return stories, subtasks, bugs


def parse_raw(rows, headers):
    col = build_column_index(headers, GENERIC_FIELD_PATTERNS)
    stories = []

    if not col:
        print("  ⚠️ Không detect được cột nào, dùng vị trí cột thô")
        for row in rows:
            values = [str_cell(v) for v in row]
            if any(values):
                stories.append({
                    "id": values[0] if values else "",
                    "summary": values[1] if len(values) > 1 else "",
                    "status": values[2] if len(values) > 2 else "To Do",
                    "assignee": "Unassigned",
                    "priority": "Medium",
                    "issue_type": "Story",
                    "start_date": None,
                    "due_date": None,
                    "task_process": "",
                    "original_estimate_h": 0,
                    "time_spent_h": 0,
                    "remaining_estimate_h": 0,
                    "parent_key": "",
                    "pct_done": 0,
                    "subtasks": [],
                })
        print(f"\n✅ Raw format: {len(stories)} stories")
        return stories, [], []

    for row in rows:
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue

        def cv(field, default=""):
            c = col.get(field)
            return str_cell(row[c]) if c is not None and c < len(row) else default

        summary = cv("summary")
        if not summary:
            continue

        status = cv("status", "To Do") or "To Do"
        estimate_h = parse_estimate_to_hours(cv("estimate"))
        time_spent_h = parse_estimate_to_hours(cv("time_spent"))

        stories.append({
            "id": cv("id") or summary[:30],
            "summary": summary[:200],
            "status": status,
            "assignee": cv("assignee", "Unassigned") or "Unassigned",
            "priority": cv("priority", "Medium") or "Medium",
            "issue_type": "Story",
            "start_date": parse_date_from_cell(row[col["start_date"]]) if "start_date" in col and col["start_date"] < len(row) else None,
            "due_date": parse_date_from_cell(row[col["due_date"]]) if "due_date" in col and col["due_date"] < len(row) else None,
            "task_process": "",
            "original_estimate_h": estimate_h,
            "time_spent_h": time_spent_h,
            "remaining_estimate_h": max(estimate_h - time_spent_h, 0),
            "parent_key": cv("parent_key"),
            "pct_done": 100 if status.lower() in ("done", "completed", "closed", "resolved") else 0,
            "subtasks": [],
        })
    print(f"\n✅ Generic format: {len(stories)} stories")
    return stories, [], []


def parse_estimate_to_hours(val):
    if not val:
        return 0
    val = str(val).strip()
    if not val:
        return 0
    try:
        num = float(val.replace(",", ""))
        if num > 1000:
            return round(num / 3600, 1)
        return num
    except ValueError:
        return 0


def parse_pct_done(val):
    if not val:
        return 0
    if isinstance(val, (int, float)):
        return float(val)
    val = str(val).strip().replace("%", "")
    try:
        return float(val)
    except ValueError:
        return 0


# ──────────────── ORCHESTRATION ────────────────

def _process_rows_to_output(
    all_rows: List[List[Any]],
    source_label: str,
    file_identifier: str,
    sheet_name: str,
    sprint_name: str,
    today: date,
    output_dir: str,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
) -> Optional[str]:
    """Hàm lõi: Phân tích các hàng dữ liệu (từ openpyxl hoặc Graph API) và xuất JSON chuẩn."""
    if len(all_rows) < 2:
        print("  → Bỏ qua (không đủ hàng dữ liệu)")
        return None

    # Tìm hàng header
    best_idx, best_score = 0, 0
    for idx, row in enumerate(all_rows):
        non_empty = sum(1 for cell in row if cell is not None and str(cell).strip())
        if non_empty >= 3:
            row_text = " ".join(str(c).lower() for c in row if c)
            score = sum(1 for kw in ["status", "assignee", "priority", "story", "task",
                                      "effort", "point", "date", "process", "pic",
                                      "module", "manday", "us id", "working day"] if kw in row_text)
            if score > best_score:
                best_score = score
                best_idx = idx

    header_idx = best_idx if best_score > 0 else 0
    headers = [str(c).strip().replace("\n", " ").replace("\r", " ") if c is not None else "" for c in all_rows[header_idx]]
    rows = all_rows[header_idx + 1:]
    rows = [r for r in rows if any(cell is not None and str(cell).strip() for cell in r)]

    print(f"  → Header row: #{header_idx + 1}")
    print(f"  → {len(rows)} data rows, {len(headers)} columns")
    print(f"  → Headers: {headers}")

    fmt = detect_format(headers)
    print(f"  → Detected format: {fmt}")

    all_stories, all_subtasks, all_bugs = [], [], []
    if fmt == "sprint_backlog":
        all_stories, all_subtasks, all_bugs = parse_sprint_backlog(rows, headers)
    elif fmt == "jira":
        all_stories, all_subtasks, all_bugs = parse_jira(rows, headers)
    else:
        all_stories, all_subtasks, all_bugs = parse_raw(rows, headers)

    output_stories = []
    if fmt == "sprint_backlog":
        output_stories = list(all_stories)
    else:
        subtask_map = {}
        for sub in all_subtasks:
            pk = sub.get("parent_key", "")
            if pk:
                subtask_map.setdefault(pk, []).append(sub)

        for story in all_stories:
            s = dict(story)
            s["subtasks"] = subtask_map.get(s["id"], [])
            output_stories.append(s)

        parent_ids = {st["id"] for st in all_stories}
        orphan_subtasks = [s for s in all_subtasks if s.get("parent_key", "") not in parent_ids]
        for sub in orphan_subtasks:
            output_stories.append({
                "id": sub["id"],
                "summary": sub["summary"],
                "status": sub["status"],
                "assignee": sub["assignee"],
                "priority": sub["priority"],
                "issue_type": "Story",
                "start_date": sub.get("start_date"),
                "due_date": sub.get("due_date"),
                "task_process": sub.get("task_process", ""),
                "original_estimate_h": sub.get("original_estimate_h", 0),
                "time_spent_h": sub.get("time_spent_h", 0),
                "remaining_estimate_h": sub.get("remaining_estimate_h", 0),
                "parent_key": "",
                "pct_done": sub.get("pct_done", 0),
                "subtasks": [],
            })

    output = {
        "source": source_label,
        "file_path": file_identifier,
        "sheet_name": sheet_name,
        "sprint_name": sprint_name or sheet_name,
        "format": fmt,
        "today": today.isoformat(),
        "exported_at": datetime.now(VN_TZ).isoformat(),
        "stories": output_stories,
        "bugs": all_bugs,
    }
    if start_date:
        output["sprint_start"] = start_date
    if end_date:
        output["sprint_end"] = end_date

    os.makedirs(output_dir, exist_ok=True)
    safe_name = (sprint_name or "Sprint").replace(" ", "_").replace("/", "_")
    out_path = os.path.join(output_dir, f"excel_data_{safe_name}.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False, default=str)

    print(f"\n✅ Đã xuất JSON: {out_path}")
    print(f"   Tổng: {len(output_stories)} stories, {sum(len(s.get('subtasks', [])) for s in output_stories)} subtasks, {len(all_bugs)} bugs")
    return out_path


def fetch_and_process_graph(
    share_url: str,
    sprint_name: str,
    today_str: str,
    output_dir: str,
    sheet_name: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
) -> Optional[str]:
    """Đọc dữ liệu bảng tính từ Microsoft Graph API (SharePoint / OneDrive sharing link) ."""
    if mg is None:
        raise ImportError("Thiếu module ms_graph_client. Vui lòng kiểm tra skills/_shared/ms_graph_client.py")

    today = date.fromisoformat(today_str) if today_str else date.today()
    print(f"🌐 Đang kết nối Microsoft Graph API cho liên kết: {share_url}")

    token = mg.get_ms_graph_token(scopes=["User.Read", "Files.ReadWrite"])
    drive_id, item_id, item_meta = mg.resolve_drive_item(share_url, token=token)
    file_name = item_meta.get("name", "sharepoint_excel.xlsx")
    print(f"  → Tệp: '{file_name}' | Drive ID: {drive_id[:15]}... | Item ID: {item_id}")

    sheets = mg.list_worksheets(drive_id, item_id, token=token)
    sheet_names = [s.get("name") for s in sheets]
    print(f"  → Danh sách worksheets: {sheet_names}")

    target_sheet = sheet_name if sheet_name and sheet_name in sheet_names else (sheet_names[0] if sheet_names else "Sheet1")
    print(f"  → Sử dụng worksheet: '{target_sheet}'")

    used_range = mg.get_worksheet_used_range(drive_id, item_id, target_sheet, token=token)
    all_rows = used_range.get("values") or []

    return _process_rows_to_output(
        all_rows=all_rows,
        source_label="excel_ms_graph",
        file_identifier=file_name,
        sheet_name=target_sheet,
        sprint_name=sprint_name,
        today=today,
        output_dir=output_dir,
        start_date=start_date,
        end_date=end_date,
    )


def fetch_and_process(file_path, sprint_name, today_str, output_dir, sheet_name=None, start_date=None, end_date=None):
    """Đọc file Excel cục bộ (.xlsx) bằng openpyxl."""
    if openpyxl is None:
        raise ImportError("Cần cài đặt openpyxl: pip install openpyxl")

    today = date.fromisoformat(today_str) if today_str else date.today()
    print(f"📖 Đang đọc file Excel cục bộ: {file_path}")
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active
        sheet_name = ws.title
    
    print(f"  → Worksheet: '{sheet_name}'")
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    return _process_rows_to_output(
        all_rows=all_rows,
        source_label="excel",
        file_identifier=file_path,
        sheet_name=sheet_name,
        sprint_name=sprint_name,
        today=today,
        output_dir=output_dir,
        start_date=start_date,
        end_date=end_date,
    )


def main():
    parser = argparse.ArgumentParser(description="Fetch Excel data (local or Microsoft Graph API) for weekly report")
    parser.add_argument("--file", help="Path to local Excel file (.xlsx) hoặc link SharePoint/OneDrive")
    parser.add_argument("--excel-url", help="Link SharePoint/OneDrive (Microsoft Graph API)")
    parser.add_argument("--share-url", help="Alias của --excel-url")
    parser.add_argument("--sprint-name", default="Current Sprint", help="Sprint name")
    parser.add_argument("--start-date", default=None, help="Start date (YYYY-MM-DD)")
    parser.add_argument("--end-date", default=None, help="End date (YYYY-MM-DD)")
    parser.add_argument("--today", default=date.today().isoformat(), help="Report date (YYYY-MM-DD)")
    parser.add_argument("--output-dir", default="/mnt/Data/ai/pm", help="Output directory")
    parser.add_argument("--sheet-name", default=None, help="Worksheet name (default: active sheet)")
    args = parser.parse_args()

    target_url = args.excel_url or args.share_url or ""
    file_arg = args.file or ""

    is_graph_url = bool(
        target_url
        or file_arg.startswith("http://")
        or file_arg.startswith("https://")
        or file_arg.startswith("u!")
        or "sharepoint.com" in file_arg
        or "1drv.ms" in file_arg
    )

    if is_graph_url:
        url_to_use = target_url or file_arg
        fetch_and_process_graph(
            share_url=url_to_use,
            sprint_name=args.sprint_name,
            today_str=args.today,
            output_dir=args.output_dir,
            sheet_name=args.sheet_name,
            start_date=args.start_date,
            end_date=args.end_date,
        )
    elif file_arg:
        fetch_and_process(
            file_path=file_arg,
            sprint_name=args.sprint_name,
            today_str=args.today,
            output_dir=args.output_dir,
            sheet_name=args.sheet_name,
            start_date=args.start_date,
            end_date=args.end_date,
        )
    else:
        print("❌ Cần cung cấp --file (.xlsx / SharePoint link) hoặc --excel-url", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
